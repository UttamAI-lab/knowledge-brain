from pathlib import Path
from docx import Document as DocxDocument
from langchain_core.documents import Document


def load_docx(file_path: str | Path) -> list[Document]:
    file_path = Path(file_path)
    doc = DocxDocument(str(file_path))

    full_text = "\n".join(
        para.text for para in doc.paragraphs if para.text.strip()
    )

    return [
        Document(
            page_content=full_text,
            metadata={
                "source": file_path.name,
                "file_path": str(file_path),
            },
        )
    ]